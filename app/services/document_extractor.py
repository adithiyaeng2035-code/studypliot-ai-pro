"""Bounded text extraction for student-submitted PDF and Excel documents."""

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 40_000
MAX_PDF_PAGES = 200
MAX_WORKSHEETS = 20
MAX_WORKSHEET_ROWS = 5_000
MAX_WORKSHEET_COLUMNS = 50
MAX_XLSX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024


class DocumentExtractionError(ValueError):
    """Raised when an uploaded document cannot be safely read as text."""


def _append_limited(parts: list[str], text: str, current_length: int) -> tuple[int, bool]:
    remaining = MAX_EXTRACTED_CHARS - current_length
    if remaining <= 0:
        return current_length, bool(text)
    if len(text) > remaining:
        parts.append(text[:remaining])
        return MAX_EXTRACTED_CHARS, True
    parts.append(text)
    return current_length + len(text), False


def _extract_pdf(data: bytes) -> dict[str, object]:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise DocumentExtractionError("This PDF is password-protected. Upload an unlocked copy.")
        page_count = len(reader.pages)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError("This PDF could not be read. Check that it is a valid PDF file.") from error

    parts: list[str] = []
    total = 0
    truncated = page_count > MAX_PDF_PAGES
    pages_to_read = min(page_count, MAX_PDF_PAGES)
    pages_with_text = 0
    for page_number in range(pages_to_read):
        try:
            text = (reader.pages[page_number].extract_text() or "").strip()
        except Exception as error:
            raise DocumentExtractionError(
                f"Text could not be extracted from PDF page {page_number + 1}."
            ) from error
        if not text:
            continue
        pages_with_text += 1
        total, page_truncated = _append_limited(
            parts, f"[PDF page {page_number + 1}]\n{text}\n", total
        )
        truncated = truncated or page_truncated
        if page_truncated:
            break

    if pages_with_text == 0:
        raise DocumentExtractionError(
            "No selectable text was found. Scanned PDFs need OCR first; this upload tool reads text-based PDFs."
        )
    if truncated:
        parts.append("\n[Document excerpt shortened to fit safe analysis limits.]\n")
    return {
        "kind": "PDF",
        "text": "\n".join(parts),
        "sources": [f"Page {number + 1}" for number in range(pages_to_read)],
        "truncated": truncated,
    }


def _extract_excel(data: bytes) -> dict[str, object]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 5_000 or sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise DocumentExtractionError("This spreadsheet is too large to analyze safely.")
    except DocumentExtractionError:
        raise
    except (BadZipFile, OSError) as error:
        raise DocumentExtractionError("This is not a valid .xlsx or .xlsm workbook.") from error

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as error:
        raise DocumentExtractionError("This Excel workbook could not be read. Upload a valid .xlsx or .xlsm file.") from error

    parts: list[str] = []
    total = 0
    truncated = len(workbook.worksheets) > MAX_WORKSHEETS
    sheet_names = [sheet.title for sheet in workbook.worksheets[:MAX_WORKSHEETS]]
    populated_rows = 0
    try:
        for worksheet in workbook.worksheets[:MAX_WORKSHEETS]:
            max_row = min(worksheet.max_row or 0, MAX_WORKSHEET_ROWS)
            max_column = min(worksheet.max_column or 0, MAX_WORKSHEET_COLUMNS)
            if (worksheet.max_row or 0) > MAX_WORKSHEET_ROWS or (worksheet.max_column or 0) > MAX_WORKSHEET_COLUMNS:
                truncated = True
            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_column,
                    values_only=False,
                ),
                start=1,
            ):
                cells = [
                    f"{cell.coordinate}={str(cell.value).strip().replace(chr(10), ' ')}"
                    for cell in row
                    if cell.value is not None and str(cell.value).strip()
                ]
                if not cells:
                    continue
                populated_rows += 1
                line = f"[Sheet: {worksheet.title}, row {row_number}] " + " | ".join(cells) + "\n"
                total, line_truncated = _append_limited(parts, line, total)
                truncated = truncated or line_truncated
                if line_truncated:
                    break
            if total >= MAX_EXTRACTED_CHARS:
                break
    finally:
        workbook.close()

    if populated_rows == 0:
        raise DocumentExtractionError("No cell values were found in this workbook.")
    if truncated:
        parts.append("\n[Workbook excerpt shortened to fit safe analysis limits.]\n")
    return {
        "kind": "Excel workbook",
        "text": "".join(parts),
        "sources": [f"Sheet: {name}" for name in sheet_names],
        "truncated": truncated,
    }


def extract_document(filename: str, data: bytes) -> dict[str, object]:
    """Extract bounded, source-labeled text from a text PDF or Excel workbook."""
    if not data:
        raise DocumentExtractionError("The uploaded file is empty.")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise DocumentExtractionError("Files must be 10 MB or smaller.")

    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension in {".xlsx", ".xlsm"}:
        return _extract_excel(data)
    raise DocumentExtractionError("Choose a text-based PDF, .xlsx, or .xlsm file.")
